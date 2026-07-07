"""报告导出工具 — 基于结构化 state 数据生成 DOCX/PDF 报告。
"""

from __future__ import annotations

import io
import os
from datetime import datetime
from typing import Any

from docx import Document
from docx.shared import Pt, RGBColor, Cm, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

_CYAN = RGBColor(0, 212, 255)
_BLACK = RGBColor(0, 0, 0)
_MUTED = RGBColor(100, 116, 139)
_DANGER = RGBColor(239, 68, 68)
_WHITE = RGBColor(226, 232, 240)
_DIM = RGBColor(71, 85, 105)


def _clean(text):
    """清洗 emoji / 变体选择符 / 不可映射字符。"""
    import re
    text = text.replace('\ufe0f', '')                    # 变体选择符
    text = re.sub(r'[\U0001f000-\U0001ffff]', '[?]', text)  # 平面 1 emoji 占位
    return text

def _r(p, text, color=None, size=None, bold=False, italic=False):
    text = _clean(text)
    run = p.add_run(text)
    run.font.name = '微软雅黑'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
    run.font.color.rgb = color or _BLACK
    if size: run.font.size = Pt(size)
    if bold: run.bold = True
    if italic: run.italic = True
    return run


def _get(state, key, default=None):
    return state.get(key, default) if isinstance(state, dict) else default


# ===================================================================
#  DOCX
# ===================================================================

def generate_docx(state: dict, session_id: str = "") -> bytes:
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(3)
    sec.bottom_margin = Cm(2.5)
    sec.left_margin = Cm(2.5)
    sec.right_margin = Cm(2.5)

    # 默认字体设置为微软雅黑 + 黑色
    style = doc.styles['Normal']
    style.font.name = '微软雅黑'
    style._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.font.size = Pt(10.5)
    style.paragraph_format.line_spacing = 1.5

    # 标题样式
    for lvl in range(1, 4):
        hs = doc.styles[f'Heading {lvl}']
        hs.font.name = '微软雅黑'
        hs._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
        hs.font.color.rgb = RGBColor(0, 0, 0)
        hs.font.bold = True

    report = state.get("report", {})
    main = _clean(report.get("main_text", ""))
    st = report.get("summary_text", {})
    impression = _clean(st.get("impression", "") if isinstance(st, dict) else str(st))

    def _add_heading(text, level=1):
        h = doc.add_heading(text, level=level)
        for r in h.runs:
            r.font.name = '微软雅黑'
            r._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
            r.font.color.rgb = RGBColor(0, 0, 0)
        return h

    def _add_body(text, size=10.5, bold=False):
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.font.name = '微软雅黑'
        run._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
        run.font.size = Pt(size)
        run.font.color.rgb = RGBColor(0, 0, 0)
        if bold:
            run.bold = True
        return p

    # ── 标题 ──
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("rare-dx · 罕见病诊断辅助报告")
    run.font.name = '微软雅黑'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
    run.font.size = Pt(18)
    run.font.color.rgb = RGBColor(0, 0, 0)
    run.bold = True

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = meta.add_run(f"生成：{datetime.now().strftime('%Y-%m-%d %H:%M')}")
    run.font.name = '微软雅黑'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(100, 116, 139)

    doc.add_paragraph("─" * 50).alignment = WD_ALIGN_PARAGRAPH.CENTER

    # 1 ── 临床印象
    if impression:
        _add_heading("一、临床印象", level=1)
        _add_body(impression, size=11)
        if report.get("llm_generated"):
            _add_body("[由 LLM 辅助生成]", size=8)

    # 2 ── 推理摘要
    if main:
        _add_heading("二、推理摘要", level=1)
        for block in main.split("\n"):
            block = block.strip()
            if not block:
                continue
            if block.startswith("【"):
                _add_body(block, size=11, bold=True)
            else:
                _add_body(block, size=10.5)

    # 3 ── 不确定性
    unc = report.get("uncertainty_notes", "")
    if unc:
        _add_heading("三、不确定性说明", level=2)
        _add_body(_clean(unc), size=10.5)

    # ── 免责声明
    doc.add_paragraph("")
    doc.add_paragraph("─" * 50).alignment = WD_ALIGN_PARAGRAPH.CENTER
    disc = doc.add_paragraph()
    disc.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = disc.add_run("本系统输出仅供临床参考，不构成诊断意见。最终诊断由接诊医师结合完整临床信息做出。")
    run.font.name = '微软雅黑'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(200, 0, 0)
    run.italic = True

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.getvalue()


# ===================================================================
#  PDF
# ===================================================================

def generate_pdf(state: dict, session_id: str = "") -> bytes:
    from fpdf import FPDF

    FONT_PATH = "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"
    has_font = os.path.exists(FONT_PATH)

    class RPDF(FPDF):
        def header(self):
            f = self._fam or "Helvetica"
            self.set_font(f, "", 8)
            self.set_text_color(100, 116, 139)
            self.cell(0, 6, _clean("rare-dx · 罕见病诊断辅助报告"), ln=True, align="C")
            self.line(self.l_margin, self.get_y(), self.w - self.r_margin, self.get_y())
            self.ln(3)
        def footer(self):
            self.set_y(-15)
            self.set_text_color(100, 116, 139)
            f = self._fam or "Helvetica"
            self.set_font(f, "", 7)
            self.cell(0, 10, f"第 {self.page_no()} 页", align="C")

    pdf = RPDF()
    pdf._fam = "CJK" if has_font else "Helvetica"
    if has_font:
        pdf.add_font("CJK", "", FONT_PATH, uni=True)
    F = pdf._fam
    pdf.set_auto_page_break(auto=True, margin=20)
    pdf.add_page()
    M = pdf.l_margin

    report = state.get("report", {})
    main = _clean(report.get("main_text", ""))
    st = report.get("summary_text", {})
    impression = _clean(st.get("impression", "") if isinstance(st, dict) else str(st))
    unc = _clean(report.get("uncertainty_notes", ""))

    def sec(txt):
        pdf.set_font(F, "", 13)
        pdf.set_text_color(0, 212, 255)
        pdf.cell(0, 10, _clean(txt), ln=True)
        pdf.ln(2)

    def body(txt, sz=10, clr=(226, 232, 240)):
        pdf.set_font(F, "", sz)
        pdf.set_text_color(*clr)
        pdf.set_x(M)
        pdf.multi_cell(0, 6, _clean(txt))
        pdf.ln(2)

    # 标题
    pdf.set_font(F, "", 18)
    pdf.set_text_color(0, 212, 255)
    pdf.cell(0, 14, _clean("rare-dx · 罕见病诊断辅助报告"), ln=True, align="C")
    pdf.set_font(F, "", 8)
    pdf.set_text_color(100, 116, 139)
    pdf.cell(0, 7, _clean(f"生成：{datetime.now().strftime('%Y-%m-%d %H:%M')}  |  会话 {session_id or '-'}"), ln=True, align="C")
    pdf.line(M, pdf.get_y() + 2, pdf.w - M, pdf.get_y() + 2)
    pdf.ln(6)

    if impression:
        sec("1. 临床印象")
        body(impression, sz=11)
    if main:
        sec("2. 推理摘要")
        for block in main.split("\n"):
            block = block.strip()
            if not block:
                continue
            if block.startswith("【"):
                pdf.set_font(F, "", 11)
                pdf.set_text_color(0, 212, 255)
                pdf.set_x(M)
                pdf.multi_cell(0, 7, _clean(block))
            else:
                body(block, sz=10)
    if unc:
        sec("3. 不确定性说明")
        body(unc, sz=10)

    # 免责
    pdf.ln(4)
    pdf.line(M, pdf.get_y(), pdf.w - M, pdf.get_y())
    pdf.ln(4)
    pdf.set_font(F, "", 9)
    pdf.set_text_color(239, 68, 68)
    pdf.multi_cell(0, 5, _clean(
        "本系统输出仅供临床参考，不构成诊断意见。最终诊断由接诊医师结合完整临床信息做出。"),
        align="C")

    raw = pdf.output(dest="S")
    if isinstance(raw, bytearray):
        return bytes(raw)
    return raw.encode("latin-1") if isinstance(raw, str) else raw
